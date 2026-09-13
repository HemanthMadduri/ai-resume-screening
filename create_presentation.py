import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    # Color Palette
    COLOR_BG = RGBColor(248, 250, 252)       # Light Slate #F8FAFC
    COLOR_HEADER = RGBColor(15, 23, 42)      # Deep Navy #0F172A
    COLOR_ACCENT = RGBColor(79, 70, 229)     # Indigo #4F46E5
    COLOR_CARD_BG = RGBColor(255, 255, 255)  # White
    COLOR_CARD_BORDER = RGBColor(226, 232, 240) # Slate border
    COLOR_TEXT = RGBColor(51, 65, 85)        # Body text #334155
    COLOR_MUTED = RGBColor(100, 116, 139)    # Muted text #64748B
    COLOR_GREEN = RGBColor(16, 185, 129)     # Emerald #10B981
    COLOR_AMBER = RGBColor(245, 158, 11)     # Amber #F59E0B
    COLOR_BLUE = RGBColor(2, 132, 199)       # Sky blue #0284C7
    COLOR_PURPLE = RGBColor(147, 51, 234)    # Violet #9333EA

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background() # No border
        return bg

    def add_slide_header(slide, number_str, title_str, subtitle_str):
        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # Pill Tag / Section
        p0 = tf.paragraphs[0]
        p0.text = f"SECTION {number_str}"
        p0.font.name = "Arial"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_ACCENT
        p0.space_after = Pt(2)
        
        # Main Title
        p1 = tf.add_paragraph()
        p1.text = title_str
        p1.font.name = "Arial"
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_HEADER
        p1.space_after = Pt(2)
        
        # Subtitle
        p2 = tf.add_paragraph()
        p2.text = subtitle_str
        p2.font.name = "Arial"
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # SLIDE 1: INTRODUCTION
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_background(slide1)
    add_slide_header(slide1, "01", "Introduction", "Intelligent AI Resume Screening and Candidate Ranking System")

    # Left Container: Key Bullet Points (Max 5 points)
    left_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.8), Inches(5.0))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = COLOR_CARD_BG
    left_card.line.color.rgb = COLOR_CARD_BORDER
    left_card.line.width = Pt(1.5)

    tf_left = left_card.text_frame
    tf_left.word_wrap = True
    tf_left.margin_left = tf_left.margin_right = tf_left.margin_top = Inches(0.3)
    
    p_title = tf_left.paragraphs[0]
    p_title.text = "Key System Highlights"
    p_title.font.name = "Arial"
    p_title.font.size = Pt(16)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_HEADER
    p_title.space_after = Pt(14)

    intro_points = [
        ("Automated Resume Screening", "Processes high volumes of applicant resumes against Job Descriptions in seconds."),
        ("Multi-Format Extraction", "Extracts candidate details from PDF, DOCX, and TXT resumes automatically."),
        ("Local NLP Engine", "Runs 100% locally with spaCy & scikit-learn without costly external AI APIs."),
        ("Transparent Ranking", "Scores candidates using weighted skills, experience, and education matching."),
        ("Explainable Decisions", "Provides human-readable reasons explaining candidate strengths and gaps.")
    ]

    for title, desc in intro_points:
        p = tf_left.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(10)

    # Right Container: Flowchart / Process Diagram
    right_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = COLOR_CARD_BG
    right_card.line.color.rgb = COLOR_CARD_BORDER
    right_card.line.width = Pt(1.5)

    tf_right = right_card.text_frame
    tf_right.word_wrap = True
    tf_right.margin_left = tf_right.margin_right = tf_right.margin_top = Inches(0.3)
    
    pr_title = tf_right.paragraphs[0]
    pr_title.text = "End-to-End System Workflow"
    pr_title.font.name = "Arial"
    pr_title.font.size = Pt(16)
    pr_title.font.bold = True
    pr_title.font.color.rgb = COLOR_HEADER
    pr_title.space_after = Pt(12)

    flow_steps = [
        ("Step 1: Input Job Description", "Recruiter selects or pastes Job Description requirements.", COLOR_ACCENT),
        ("Step 2: Bulk Resume Upload", "Uploads multiple PDF / DOCX / TXT candidate resumes.", COLOR_BLUE),
        ("Step 3: NLP Extraction & Cleaning", "Extracts entities, skills, experience duration, and education.", COLOR_PURPLE),
        ("Step 4: Multi-Factor Scoring & Ranking", "Computes weighted qualification score (0-100%) and ranks.", COLOR_GREEN),
        ("Step 5: Explainable Recruiter Review", "Displays matched/missing skills, profile inspection, and export.", COLOR_AMBER)
    ]

    for i, (stitle, sdesc, scolor) in enumerate(flow_steps):
        step_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(2.5 + (i * 0.8)), Inches(5.1), Inches(0.68))
        step_box.fill.solid()
        step_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
        step_box.line.color.rgb = scolor
        step_box.line.width = Pt(1.5)
        
        stf = step_box.text_frame
        stf.word_wrap = True
        stf.margin_left = Inches(0.15)
        stf.margin_top = Inches(0.08)
        
        sp1 = stf.paragraphs[0]
        sp1.text = stitle
        sp1.font.name = "Arial"
        sp1.font.size = Pt(11)
        sp1.font.bold = True
        sp1.font.color.rgb = scolor
        
        sp2 = stf.add_paragraph()
        sp2.text = sdesc
        sp2.font.name = "Arial"
        sp2.font.size = Pt(9.5)
        sp2.font.color.rgb = COLOR_TEXT

    # =========================================================================
    # SLIDE 2: BACKGROUND KNOWLEDGE
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_background(slide2)
    add_slide_header(slide2, "02", "Background Knowledge", "Core Technologies & Concepts Powering the ATS Engine")

    # Left Card: Summary Bullet Points
    left_card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.8), Inches(5.0))
    left_card2.fill.solid()
    left_card2.fill.fore_color.rgb = COLOR_CARD_BG
    left_card2.line.color.rgb = COLOR_CARD_BORDER
    left_card2.line.width = Pt(1.5)

    tf2 = left_card2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = tf2.margin_top = Inches(0.3)
    
    p2_title = tf2.paragraphs[0]
    p2_title.text = "Key Theoretical Concepts"
    p2_title.font.name = "Arial"
    p2_title.font.size = Pt(16)
    p2_title.font.bold = True
    p2_title.font.color.rgb = COLOR_HEADER
    p2_title.space_after = Pt(14)

    bg_points = [
        ("Natural Language Processing (NLP)", "Enables computers to understand and extract structured knowledge from unstructured resume text."),
        ("spaCy Named Entity Recognition", "Identifies entities such as person names, contact information, organizations, and locations."),
        ("TF-IDF & Cosine Similarity", "Measures the contextual and lexical similarity score between resumes and Job Descriptions."),
        ("Canonical Skill Taxonomy", "Normalizes synonyms and acronyms (e.g. ML = Machine Learning, k8s = Kubernetes)."),
        ("Multi-Tier Document Parsers", "Utilizes pdfplumber, python-docx, and regex to read diverse resume formats reliably.")
    ]

    for title, desc in bg_points:
        p = tf2.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(10)

    # Right Card: 4 Technology Architecture Blocks
    tech_blocks = [
        ("spaCy (en_core_web_sm)", "Industrial-strength NLP library used for POS tagging, lemmatization, and Named Entity Recognition.", COLOR_ACCENT, Inches(6.8), Inches(1.8)),
        ("Scikit-Learn (TF-IDF)", "Computes term frequency-inverse document frequency vectors and cosine similarity index.", COLOR_BLUE, Inches(9.8), Inches(1.8)),
        ("Skill & Alias Graph", "Dictionary taxonomy connecting technical variations (React.js, React, ReactJS) to single canonical skills.", COLOR_GREEN, Inches(6.8), Inches(4.35)),
        ("Python & Flask Backend", "Lightweight, responsive web framework serving RESTful APIs and modern responsive ATS views.", COLOR_PURPLE, Inches(9.8), Inches(4.35))
    ]

    for title, desc, col, x, y in tech_blocks:
        tb = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.75), Inches(2.45))
        tb.fill.solid()
        tb.fill.fore_color.rgb = COLOR_CARD_BG
        tb.line.color.rgb = col
        tb.line.width = Pt(1.5)
        
        tbf = tb.text_frame
        tbf.word_wrap = True
        tbf.margin_left = tbf.margin_right = tbf.margin_top = Inches(0.2)
        
        tp1 = tbf.paragraphs[0]
        tp1.text = title
        tp1.font.name = "Arial"
        tp1.font.size = Pt(13)
        tp1.font.bold = True
        tp1.font.color.rgb = col
        tp1.space_after = Pt(8)
        
        tp2 = tbf.add_paragraph()
        tp2.text = desc
        tp2.font.name = "Arial"
        tp2.font.size = Pt(11)
        tp2.font.color.rgb = COLOR_TEXT

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_background(slide3)
    add_slide_header(slide3, "03", "Problem Statement", "Limitations of Manual Resume Screening vs. Our AI Solution")

    # Left Card: Core Problem Bullets (Max 5 points)
    left_card3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.8), Inches(5.0))
    left_card3.fill.solid()
    left_card3.fill.fore_color.rgb = COLOR_CARD_BG
    left_card3.line.color.rgb = COLOR_CARD_BORDER
    left_card3.line.width = Pt(1.5)

    tf3 = left_card3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_right = tf3.margin_top = Inches(0.25)
    
    p3_title = tf3.paragraphs[0]
    p3_title.text = "Challenges in Recruitment"
    p3_title.font.name = "Arial"
    p3_title.font.size = Pt(15)
    p3_title.font.bold = True
    p3_title.font.color.rgb = COLOR_HEADER
    p3_title.space_after = Pt(12)

    prob_points = [
        ("Massive Volume", "Recruiters receive hundreds of resumes per job posting, creating bottlenecks."),
        ("Time Constraints", "Manual review takes only 6–7 seconds per resume, causing high error rates."),
        ("Human & Demographic Bias", "Unconscious bias affects candidate selection based on age, gender, or name."),
        ("Keyword Search Flaws", "Basic keyword search misses qualified candidates using skill synonyms (e.g. PyTorch vs Torch)."),
        ("Lack of Explainability", "Existing AI hiring tools act as black boxes with no justification for rankings.")
    ]

    for title, desc in prob_points:
        p = tf3.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(8)

    # Right Card: Comparison Table (Manual vs AI ATS)
    table_shape = slide3.shapes.add_table(6, 3, Inches(5.8), Inches(1.8), Inches(6.7), Inches(5.0))
    table = table_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(2.4)
    table.columns[2].width = Inches(2.5)

    headers = ["Criterion", "Traditional Manual Screening", "Our AI Resume Screening System"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER
        for p in cell.text_frame.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER

    table_data = [
        ("Screening Speed", "Several hours / days per batch", "Instant (~2-5 seconds for 100+ resumes)"),
        ("Skill Recognition", "Strict exact-match keyword searching", "Intelligent alias & synonym graph matching"),
        ("Bias & Fairness", "Prone to unconscious human bias", "100% merit-based (Ignores demographics)"),
        ("Scoring Accuracy", "Subjective and inconsistent ratings", "Transparent weighted formula (Skills, Exp, Edu)"),
        ("Decision Output", "No explanations provided", "Clear explainability (strengths & missing gaps)")
    ]

    for row_idx, row in enumerate(table_data, start=1):
        bg_col = RGBColor(255, 255, 255) if row_idx % 2 == 1 else RGBColor(241, 245, 249)
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_col
            for p in cell.text_frame.paragraphs:
                p.font.name = "Arial"
                p.font.size = Pt(10)
                p.font.color.rgb = COLOR_HEADER if col_idx == 0 else (COLOR_ACCENT if col_idx == 2 else COLOR_TEXT)
                if col_idx == 0:
                    p.font.bold = True

    # =========================================================================
    # SLIDE 4: LITERATURE SURVEY
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_background(slide4)
    add_slide_header(slide4, "04", "Literature Survey", "Analysis of Existing Approaches vs. Our Proposed System")

    # Full Width Comparison Table for Literature Survey
    table_shape4 = slide4.shapes.add_table(5, 4, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.9))
    table4 = table_shape4.table
    table4.columns[0].width = Inches(2.2)
    table4.columns[1].width = Inches(3.1)
    table4.columns[2].width = Inches(3.2)
    table4.columns[3].width = Inches(3.233)

    lit_headers = ["Existing Approach", "Methodology / Techniques", "Identified Limitations", "Our Proposed Solution"]
    for col_idx, h in enumerate(lit_headers):
        cell = table4.cell(0, col_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER
        for p in cell.text_frame.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(11.5)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER

    lit_data = [
        ("Keyword String Search", "Regular expressions and literal string lookup", "Misses acronyms, abbreviations, and related technologies; rigid format constraints.", "Dynamic skill graph mapping 100+ synonyms & aliases (e.g. ML = Machine Learning)."),
        ("Basic Vector Space Models", "Word count and simple TF-IDF representations", "Ignores qualification factors like experience duration, degree levels, and certifications.", "Multi-factor weighted scoring (Skills 40%, Exp 25%, Edu 15%, Similarity 10%, Certs 10%)."),
        ("Deep Learning / Black-Box Models", "Complex neural networks (BERT / CNNs)", "High computational overhead; cannot explain why a candidate was ranked high or rejected.", "Explainable AI (XAI) bullet points showing exact matched skills and identified gaps."),
        ("Cloud / Paid AI API Platforms", "Third-party commercial APIs (OpenAI / Claude)", "High recurring subscription costs; serious candidate PII privacy & data security concerns.", "100% local execution using open-source spaCy & scikit-learn on standard hardware.")
    ]

    for row_idx, row in enumerate(lit_data, start=1):
        bg_col = RGBColor(255, 255, 255) if row_idx % 2 == 1 else RGBColor(241, 245, 249)
        for col_idx, text in enumerate(row):
            cell = table4.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_col
            for p in cell.text_frame.paragraphs:
                p.font.name = "Arial"
                p.font.size = Pt(10.5)
                p.font.color.rgb = COLOR_HEADER if col_idx == 0 else (COLOR_ACCENT if col_idx == 3 else COLOR_TEXT)
                if col_idx == 0 or col_idx == 3:
                    p.font.bold = True

    # =========================================================================
    # SLIDE 5: SCOPE AND OBJECTIVES
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_background(slide5)
    add_slide_header(slide5, "05", "Scope and Objectives", "Project Goals, Boundaries, and Transparent Scoring Model")

    # Left Card: Project Objectives (Max 5 points)
    left_card5 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.8), Inches(5.0))
    left_card5.fill.solid()
    left_card5.fill.fore_color.rgb = COLOR_CARD_BG
    left_card5.line.color.rgb = COLOR_CARD_BORDER
    left_card5.line.width = Pt(1.5)

    tf5 = left_card5.text_frame
    tf5.word_wrap = True
    tf5.margin_left = tf5.margin_right = tf5.margin_top = Inches(0.3)
    
    p5_title = tf5.paragraphs[0]
    p5_title.text = "Key Project Objectives"
    p5_title.font.name = "Arial"
    p5_title.font.size = Pt(16)
    p5_title.font.bold = True
    p5_title.font.color.rgb = COLOR_HEADER
    p5_title.space_after = Pt(14)

    objectives = [
        ("Multi-Format Document Ingestion", "Accurately parse PDF, DOCX, and TXT resumes in bulk with resilient error handling."),
        ("Skill & Experience Alignment", "Extract technical skills and experience duration, comparing them against JD requirements."),
        ("Transparent Weighted Scoring", "Implement an explainable scoring formula evaluating 5 core professional criteria."),
        ("Demographic Fairness", "Ensure zero bias by completely ignoring age, gender, race, and personal attributes."),
        ("Interactive Recruiter Dashboard", "Provide a modern web interface with search, filtering, shortlisting, and CSV export.")
    ]

    for title, desc in objectives:
        p = tf5.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(10)

    # Right Card: Weighted Scoring Model Visual Breakdown
    right_card5 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    right_card5.fill.solid()
    right_card5.fill.fore_color.rgb = COLOR_CARD_BG
    right_card5.line.color.rgb = COLOR_CARD_BORDER
    right_card5.line.width = Pt(1.5)

    tf5_right = right_card5.text_frame
    tf5_right.word_wrap = True
    tf5_right.margin_left = tf5_right.margin_right = tf5_right.margin_top = Inches(0.3)
    
    p5r_title = tf5_right.paragraphs[0]
    p5r_title.text = "Configurable Scoring Formula (100%)"
    p5r_title.font.name = "Arial"
    p5r_title.font.size = Pt(16)
    p5r_title.font.bold = True
    p5r_title.font.color.rgb = COLOR_HEADER
    p5r_title.space_after = Pt(12)

    weights_bars = [
        ("Technical Skills Match (Required + Preferred)", "40%", "40%", COLOR_ACCENT, Inches(2.6)),
        ("Relevant Work Experience Duration", "25%", "25%", COLOR_BLUE, Inches(3.35)),
        ("Education Qualification Level", "15%", "15%", COLOR_PURPLE, Inches(4.1)),
        ("Semantic JD Cosine Similarity (TF-IDF)", "10%", "10%", COLOR_GREEN, Inches(4.85)),
        ("Certifications & Verified Qualifications", "10%", "10%", COLOR_AMBER, Inches(5.6))
    ]

    for label, weight_label, pct, col, y_pos in weights_bars:
        bar_bg = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), y_pos, Inches(5.1), Inches(0.62))
        bar_bg.fill.solid()
        bar_bg.fill.fore_color.rgb = RGBColor(241, 245, 249)
        bar_bg.line.color.rgb = COLOR_CARD_BORDER
        bar_bg.line.width = Pt(1)
        
        btf = bar_bg.text_frame
        btf.word_wrap = True
        btf.margin_left = Inches(0.15)
        btf.margin_top = Inches(0.06)
        
        bp1 = btf.paragraphs[0]
        bp1.text = label
        bp1.font.name = "Arial"
        bp1.font.size = Pt(10.5)
        bp1.font.bold = True
        bp1.font.color.rgb = COLOR_HEADER
        
        bp2 = btf.add_paragraph()
        bp2.text = f"Weight Contribution: {weight_label}"
        bp2.font.name = "Arial"
        bp2.font.size = Pt(9.5)
        bp2.font.bold = True
        bp2.font.color.rgb = col

    # Save Presentation
    output_path = os.path.join(os.path.dirname(__file__), "AI_Resume_Screening_Presentation.pptx")
    prs.save(output_path)
    
    # Also save to parent directory
    parent_output = os.path.join("C:\\mini project", "AI_Resume_Screening_Presentation.pptx")
    try:
        prs.save(parent_output)
    except Exception:
        pass
        
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == "__main__":
    create_deck()
