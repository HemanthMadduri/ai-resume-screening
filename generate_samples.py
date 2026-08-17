import os
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
SAMPLE_JDS_DIR = os.path.join(BASE_DIR, "data", "samples", "sample_jds")
SAMPLE_RESUMES_DIR = os.path.join(BASE_DIR, "data", "samples", "sample_resumes")

os.makedirs(SAMPLE_JDS_DIR, exist_ok=True)
os.makedirs(SAMPLE_RESUMES_DIR, exist_ok=True)

# -------------------------------------------------------------
# 1. SAMPLE JOB DESCRIPTIONS
# -------------------------------------------------------------
JD_FULLSTACK = """Job Title: Senior Full Stack Engineer
Location: San Francisco, CA / Remote
Experience Required: 4+ years of experience
Education: Bachelor's degree in Computer Science, Software Engineering, or related field

Job Overview:
We are looking for a highly skilled Senior Full Stack Engineer to design, develop, and scale our core SaaS platform. You will collaborate with cross-functional product teams to deliver high-performance, fault-tolerant web applications.

Required Skills:
- Python, Django or FastAPI
- React, JavaScript, TypeScript, HTML5, CSS3
- PostgreSQL, Redis, Database Optimization
- REST API design, Microservices architecture
- Docker, CI/CD, Git

Preferred Skills:
- AWS, Kubernetes, Terraform
- GraphQL, WebSockets
- Automated Unit Testing, TDD, Agile/Scrum

Key Responsibilities:
- Build responsive front-end user interfaces using React and modern JavaScript.
- Develop robust, scalable back-end APIs in Python.
- Architect high-throughput database schemas with PostgreSQL.
- Implement CI/CD deployment pipelines using Docker and GitHub Actions.
"""

JD_ML = """Job Title: Machine Learning & NLP Specialist
Location: New York, NY / Remote
Experience Required: 3+ years of experience in ML / Data Science
Education: Master's degree or Bachelor's in Computer Science, Data Science, or AI

Job Overview:
We are seeking an experienced Machine Learning & NLP Engineer to lead the development of our intelligent text extraction, classification, and ranking systems.

Required Skills:
- Python, PyTorch or TensorFlow, scikit-learn
- Natural Language Processing (NLP), spaCy, Hugging Face, Transformers
- Machine Learning algorithms, Data Analysis, Pandas, NumPy
- REST API development using FastAPI or Flask
- Docker, Git, Linux

Preferred Skills:
- LLMs, Generative AI, Computer Vision (OpenCV)
- AWS or Google Cloud Platform (GCP)
- MLflow, MLOps, Apache Spark
- DeepLearning.AI Specialization or relevant ML certifications

Key Responsibilities:
- Design and deploy state-of-the-art NLP and transformer models.
- Build automated data processing pipelines with Pandas and PyTorch.
- Deploy scalable ML inference APIs using FastAPI and Docker.
- Collaborate with software engineers to integrate AI models into production.
"""

with open(os.path.join(SAMPLE_JDS_DIR, "senior_fullstack_engineer.txt"), "w", encoding="utf-8") as f:
    f.write(JD_FULLSTACK)

with open(os.path.join(SAMPLE_JDS_DIR, "machine_learning_engineer.txt"), "w", encoding="utf-8") as f:
    f.write(JD_ML)

print("Sample JDs created.")

# -------------------------------------------------------------
# 2. SAMPLE RESUMES DATA
# -------------------------------------------------------------
CANDIDATES_DATA = [
    {
        "filename": "rahul_sharma_senior_ml.pdf",
        "format": "pdf",
        "name": "Rahul Sharma",
        "email": "rahul.sharma@example.com",
        "phone": "+1 415-555-0142",
        "location": "San Francisco, CA",
        "title": "Senior Machine Learning & NLP Engineer",
        "summary": "Passionate Machine Learning Engineer with 5+ years of experience specializing in Natural Language Processing (NLP), LLMs, transformer models, and scalable AI systems.",
        "skills": "Python, PyTorch, scikit-learn, spaCy, Hugging Face, Natural Language Processing, Machine Learning, Deep Learning, Pandas, NumPy, FastAPI, Docker, Git, Linux, AWS, LLMs",
        "experience": [
            {
                "role": "Lead ML Engineer",
                "company": "NeuralTech AI Solutions",
                "period": "2021 - Present",
                "details": "Architected end-to-end NLP pipelines using Hugging Face and PyTorch processing 5M+ daily documents. Deployed low-latency inference microservices using FastAPI and Docker on AWS. Mentored 4 junior ML engineers."
            },
            {
                "role": "Data Scientist & NLP Specialist",
                "company": "DataMatrix Labs",
                "period": "2019 - 2021",
                "details": "Developed custom named entity recognition (NER) models with spaCy and scikit-learn. Engineered data processing workflows using Pandas and NumPy."
            }
        ],
        "education": "Master of Science in Computer Science & AI, Stanford University (2017 - 2019)\nBachelor of Technology in Computer Science, IIT Bombay (2013 - 2017)",
        "certifications": "DeepLearning.AI Specialization, AWS Certified Cloud Practitioner"
    },
    {
        "filename": "priya_kumar_fullstack.docx",
        "format": "docx",
        "name": "Priya Kumar",
        "email": "priya.kumar@example.com",
        "phone": "+1 206-555-0189",
        "location": "Seattle, WA",
        "title": "Staff Full Stack Engineer",
        "summary": "Senior Full Stack Developer with 6 years of proven experience building scalable web applications, microservices, and reactive frontends. Expert in React, Python, Django, and cloud deployment.",
        "skills": "Python, Django, FastAPI, React, JavaScript, TypeScript, HTML5, CSS3, PostgreSQL, Redis, REST API, Microservices, Docker, Kubernetes, AWS, Git, CI/CD, Agile",
        "experience": [
            {
                "role": "Senior Full Stack Engineer",
                "company": "CloudScale Platforms",
                "period": "2020 - Present",
                "details": "Designed and implemented scalable SaaS dashboard using React, TypeScript, and Python Django back-end. Managed PostgreSQL databases and Redis caching layer, boosting API throughput by 45%. Built automated CI/CD pipelines with GitHub Actions."
            },
            {
                "role": "Full Stack Software Developer",
                "company": "WebFront Systems",
                "period": "2018 - 2020",
                "details": "Developed responsive user interfaces with React and Redux. Built RESTful APIs using Python and Flask with PostgreSQL database backends."
            }
        ],
        "education": "Bachelor of Science in Software Engineering, University of Washington (2014 - 2018)",
        "certifications": "AWS Certified Developer, Certified Scrum Master"
    },
    {
        "filename": "arjun_patel_backend.txt",
        "format": "txt",
        "name": "Arjun Patel",
        "email": "arjun.patel@example.com",
        "phone": "+1 512-555-0177",
        "location": "Austin, TX",
        "title": "Backend Python Developer",
        "summary": "Backend Software Engineer with 3.5 years of experience building high-throughput APIs, distributed systems, and database architectures.",
        "skills": "Python, FastAPI, Django, PostgreSQL, Redis, Docker, Git, REST API, Linux, SQL, MongoDB, Unit Testing, CI/CD",
        "experience": [
            {
                "role": "Backend Engineer",
                "company": "Apex Fintech Systems",
                "period": "2021 - Present",
                "details": "Engineered transactional financial APIs in Python using FastAPI and PostgreSQL. Implemented Redis caching and Celery task queues. Conducted rigorous unit testing and automated integration testing."
            },
            {
                "role": "Junior Backend Developer",
                "company": "ByteCraft Technologies",
                "period": "2020 - 2021",
                "details": "Assisted in migrating legacy monolith services to Python REST API microservices with Docker."
            }
        ],
        "education": "Bachelor of Technology in Information Technology, UT Austin (2016 - 2020)",
        "certifications": "AWS Certified Cloud Practitioner"
    },
    {
        "filename": "sneha_rao_junior_dev.pdf",
        "format": "pdf",
        "name": "Sneha Rao",
        "email": "sneha.rao@example.com",
        "phone": "+1 617-555-0133",
        "location": "Boston, MA",
        "title": "Junior Frontend & Web Developer",
        "summary": "Enthusiastic Junior Developer with 1 year of professional experience in frontend web development, responsive UI design, and JavaScript.",
        "skills": "JavaScript, HTML5, CSS3, React, Git, Bootstrap, Tailwind CSS, REST API, Problem Solving",
        "experience": [
            {
                "role": "Associate Web Developer",
                "company": "PixelCraft Studios",
                "period": "2023 - Present",
                "details": "Built reusable React component library with HTML5 and modern CSS3. Integrated third-party RESTful APIs and ensured cross-browser compatibility across mobile and desktop devices."
            }
        ],
        "education": "Bachelor of Science in Computer Science, Boston University (2019 - 2023)",
        "certifications": "Meta Front-End Developer Certificate"
    },
    {
        "filename": "vikram_singh_devops.docx",
        "format": "docx",
        "name": "Vikram Singh",
        "email": "vikram.singh@example.com",
        "phone": "+1 408-555-0199",
        "location": "San Jose, CA",
        "title": "DevOps & Cloud Infrastructure Engineer",
        "summary": "DevOps Specialist with 4 years of experience architecting cloud infrastructure on AWS and Kubernetes. Proven expertise in Terraform, CI/CD automation, and Linux administration.",
        "skills": "Kubernetes, Docker, AWS, Terraform, Ansible, Jenkins, CI/CD, Linux, Python, Bash/Shell, Prometheus, Grafana, Microservices, Git",
        "experience": [
            {
                "role": "Senior Cloud & DevOps Engineer",
                "company": "InfraScale Technologies",
                "period": "2022 - Present",
                "details": "Managed 15+ production Kubernetes (EKS) clusters on AWS. Automated multi-environment infrastructure provisioning with Terraform. Implemented comprehensive monitoring using Prometheus and Grafana."
            },
            {
                "role": "DevOps Engineer",
                "company": "CloudNet Solutions",
                "period": "2020 - 2022",
                "details": "Built CI/CD deployment pipelines using Jenkins and GitHub Actions. Automated Linux server maintenance using Python and Bash scripting."
            }
        ],
        "education": "Bachelor of Engineering in Computer Science, San Jose State University (2016 - 2020)",
        "certifications": "Certified Kubernetes Administrator (CKA), AWS Certified Solutions Architect"
    }
]

# Helper for PDF generation using ReportLab
def create_pdf_resume(cand: dict, filepath: str):
    doc = SimpleDocTemplate(filepath, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'CandName',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'CandSub',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#475569"),
        spaceAfter=12
    )
    section_heading = ParagraphStyle(
        'SecHead',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=10,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#334155"),
        spaceAfter=6
    )
    
    story = []
    # Header
    story.append(Paragraph(cand["name"], title_style))
    contact_line = f"{cand['title']} | {cand['email']} | {cand['phone']} | {cand['location']}"
    story.append(Paragraph(contact_line, subtitle_style))
    story.append(Spacer(1, 4))
    
    # Summary
    story.append(Paragraph("<b>PROFESSIONAL SUMMARY</b>", section_heading))
    story.append(Paragraph(cand["summary"], body_style))
    
    # Skills
    story.append(Paragraph("<b>TECHNICAL SKILLS</b>", section_heading))
    story.append(Paragraph(cand["skills"], body_style))
    
    # Experience
    story.append(Paragraph("<b>WORK EXPERIENCE</b>", section_heading))
    for exp in cand["experience"]:
        exp_header = f"<b>{exp['role']}</b> – {exp['company']} (<i>{exp['period']}</i>)"
        story.append(Paragraph(exp_header, body_style))
        story.append(Paragraph(exp["details"], body_style))
        story.append(Spacer(1, 2))
        
    # Education
    story.append(Paragraph("<b>EDUCATION</b>", section_heading))
    story.append(Paragraph(cand["education"].replace("\n", "<br/>"), body_style))
    
    # Certifications
    if cand.get("certifications"):
        story.append(Paragraph("<b>CERTIFICATIONS</b>", section_heading))
        story.append(Paragraph(cand["certifications"], body_style))
        
    doc.build(story)

# Helper for DOCX generation using python-docx
def create_docx_resume(cand: dict, filepath: str):
    doc = docx.Document()
    
    # Title
    h1 = doc.add_heading(cand["name"], level=0)
    
    # Subtitle / Contact
    p_sub = doc.add_paragraph(f"{cand['title']} | {cand['email']} | {cand['phone']} | {cand['location']}")
    
    # Summary
    doc.add_heading("PROFESSIONAL SUMMARY", level=1)
    doc.add_paragraph(cand["summary"])
    
    # Skills
    doc.add_heading("TECHNICAL SKILLS", level=1)
    doc.add_paragraph(cand["skills"])
    
    # Experience
    doc.add_heading("WORK EXPERIENCE", level=1)
    for exp in cand["experience"]:
        p_exp = doc.add_paragraph()
        r_role = p_exp.add_run(f"{exp['role']} – {exp['company']}")
        r_role.bold = True
        p_exp.add_run(f" ({exp['period']})\n")
        p_exp.add_run(exp["details"])
        
    # Education
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph(cand["education"])
    
    # Certifications
    if cand.get("certifications"):
        doc.add_heading("CERTIFICATIONS", level=1)
        doc.add_paragraph(cand["certifications"])
        
    doc.save(filepath)

# Helper for TXT generation
def create_txt_resume(cand: dict, filepath: str):
    content = f"""{cand['name']}
{cand['title']}
Email: {cand['email']} | Phone: {cand['phone']} | Location: {cand['location']}

==================================================
PROFESSIONAL SUMMARY
==================================================
{cand['summary']}

==================================================
TECHNICAL SKILLS
==================================================
{cand['skills']}

==================================================
WORK EXPERIENCE
==================================================
"""
    for exp in cand["experience"]:
        content += f"\n{exp['role']} - {exp['company']} ({exp['period']})\n{exp['details']}\n"

    content += f"""
==================================================
EDUCATION
==================================================
{cand['education']}

==================================================
CERTIFICATIONS
==================================================
{cand.get('certifications', 'None')}
"""
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content.strip())


for cand in CANDIDATES_DATA:
    target_path = os.path.join(SAMPLE_RESUMES_DIR, cand["filename"])
    if cand["format"] == "pdf":
        create_pdf_resume(cand, target_path)
    elif cand["format"] == "docx":
        create_docx_resume(cand, target_path)
    elif cand["format"] == "txt":
        create_txt_resume(cand, target_path)
    print(f"Generated sample resume: {cand['filename']}")

print("All sample resumes and JDs generated successfully!")
